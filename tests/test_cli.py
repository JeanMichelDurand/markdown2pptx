"""The command line, run in-process on real files."""
import pytest
from pptx import Presentation

from markdown2pptx import __version__
from markdown2pptx.cli import convert_files, main
from helpers import png, template

MD = "# Talk\n\n## One\n\ntext\n\n![chart](img/c.png)\n"


@pytest.fixture
def talk(tmp_path):
    (tmp_path / "img").mkdir()
    (tmp_path / "img" / "c.png").write_bytes(png())
    path = tmp_path / "talk.md"
    path.write_text(MD, encoding="utf-8")
    return path


def test_converts_next_to_the_input_with_images_found_beside_it(talk, capsys):
    assert main([str(talk)]) == 0
    out = talk.with_suffix(".pptx")
    assert capsys.readouterr().out == f"{out}: 1 content slide (1 picture), 0 chapters\n"
    assert Presentation(out).slides[0].shapes.title.text == "Talk"


def test_template_option_and_environment(talk, tmp_path, monkeypatch, capsys):
    tpl = tmp_path / "brand.pptx"
    tpl.write_bytes(template())
    assert main([str(talk), "--template", str(tpl), "--aspect", "16:9", "-o", str(tmp_path / "a.pptx")]) == 0
    assert "--aspect is ignored with --template" in capsys.readouterr().err
    assert Presentation(tmp_path / "a.pptx").slide_width == 9144000
    monkeypatch.setenv("MARKDOWN2PPTX_TEMPLATE", str(tpl))
    assert main([str(talk), "-o", str(tmp_path / "b.pptx")]) == 0
    assert Presentation(tmp_path / "b.pptx").slide_width == 9144000
    assert main([str(talk), "--template", "", "-o", str(tmp_path / "c.pptx")]) == 0     # '' turns it off
    assert Presentation(tmp_path / "c.pptx").slide_width == 12192000


def test_errors(talk, tmp_path, capsys):
    (tmp_path / "bad.pptx").write_bytes(b"nope")
    assert main([str(talk), "--template", str(tmp_path / "bad.pptx")]) == 1
    assert capsys.readouterr().err.strip() == f"error: {tmp_path / 'bad.pptx'}: not a PowerPoint .pptx or .potx file"
    assert main([str(tmp_path / "missing.md")]) == 1
    assert "cannot read" in capsys.readouterr().err
    assert main([str(talk), "--mermaid-color", "nope"]) == 1
    assert "unknown colour 'nope'" in capsys.readouterr().err


def test_batch_conversion_for_the_windows_dialog(talk, tmp_path):
    ok, report = convert_files([str(talk), str(tmp_path / "missing.md")])
    assert not ok and "1 content slide" in report and "cannot read" in report


def test_version(capsys):
    with pytest.raises(SystemExit) as e:
        main(["--version"])
    assert e.value.code == 0
    assert capsys.readouterr().out.strip() == f"markdown2pptx {__version__}"
