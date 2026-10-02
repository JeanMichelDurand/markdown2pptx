"""The drawn deck: slides, placeholders, text, native diagrams, tables, pictures, notes."""
from pptx.oxml.ns import qn
from pptx.util import Pt

from markdown2pptx import Options, convert
from helpers import EXAMPLES, layout_problems, png, reopen, shapes

FLOW = "```mermaid\nflowchart LR\n  A[Start] --> B{Check}\n  B -->|ok| C[Done]\n  B -->|no| A\n```\n"


def test_title_contents_chapters_and_slides():
    prs, deck = convert("# Talk\n\nSub\n\n## A\n### one\nx\n## B\n### two\ny\n")
    prs = reopen(prs)
    titles = [s.shapes.title.text if s.shapes.title else "" for s in prs.slides]
    assert titles == ["Talk", "Contents", "A", "one", "B", "two"]
    assert [sh.text_frame.text for sh in shapes(prs.slides[0], "deck:subtitle")] == ["Sub"]
    row = next(sh for sh in prs.slides[1].shapes if sh.has_text_frame and sh.text_frame.text == "A")
    assert row.click_action.target_slide == prs.slides[2]                           # contents rows link
    assert prs.core_properties.title == "Talk"


def test_text_goes_in_the_body_placeholder_with_runs_and_bullets():
    prs = reopen(convert("## S\n\nSome **bold** and [a link](https://e.org).\n\n- one\n  - two\n1. first\n")[0])
    body = shapes(prs.slides[1], "deck:body")[0]
    assert body.is_placeholder
    paras = body.text_frame.paragraphs
    assert [p.text for p in paras] == ["Some bold and a link.", "one", "two", "first"]
    assert paras[0].runs[1].font.bold and paras[0].runs[3].hyperlink.address == "https://e.org"
    assert paras[0]._p.pPr.find(qn("a:buNone")) is not None
    assert [p.level for p in paras[1:3]] == [0, 1]
    assert paras[3]._p.pPr.find(qn("a:buAutoNum")).get("type") == "arabicPeriod"


def test_long_text_gets_smaller_and_too_long_text_warns():
    short = reopen(convert("## S\n\nOne line.\n")[0])
    long_ = reopen(convert("## S\n\n" + "\n".join(f"- point number {i} with some words" for i in range(14)))[0])
    size = lambda prs: shapes(prs.slides[1], "deck:body")[0].text_frame.paragraphs[0].runs[0].font.size   # noqa: E731
    assert size(short) == Pt(24) and size(long_) < Pt(24)
    _, deck = convert("## S\n\n" + "\n".join(f"- item {i} " + "word " * 30 for i in range(30)))
    assert deck.warnings == ["slide 2 (S): its text does not fit: split it with a --- line"]


def test_a_mermaid_block_is_a_native_group_with_glued_connectors():
    prs = reopen(convert("## Flow\n\n" + FLOW)[0])
    slide = prs.slides[1]
    (group,) = shapes(slide, "deck:diagram")
    ids = [int(c.get("id")) for c in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    assert len(ids) == len(set(ids))                                      # unique on the slide
    inner = {c.get("id") for c in group._element.iter(qn("p:cNvPr"))}
    ends = [c.get("id") for tag in ("a:stCxn", "a:endCxn") for c in group._element.iter(qn(tag))]
    assert ends and set(ends) <= inner                                    # glued to shapes of the group
    texts = {sh.text_frame.text for sh in group.shapes if sh.has_text_frame and sh.text_frame.text}
    assert {"Start", "Check", "Done"} <= texts
    assert 'schemeClr val="accent1"' in group._element.xml                # theme colours: the template's


def test_diagram_fits_its_box_and_fonts_scale_with_it():
    prs = reopen(convert("## Flow\n\n" + FLOW)[0])
    assert layout_problems(prs) == []
    group = shapes(prs.slides[1], "deck:diagram")[0]
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    r = int(xfrm.find(qn("a:ext")).get("cx")) / int(xfrm.find(qn("a:chExt")).get("cx"))
    sizes = {int(e.get("sz")) for e in group._element.iter(qn("a:rPr")) if e.get("sz")}
    assert r > 1 and max(sizes) == round(1200 * r / 50) * 50             # 12 pt drawn, scaled like the shapes


def test_a_diagram_that_cannot_be_drawn_shows_its_source():
    prs, deck = convert("## Pie\n\n```mermaid\npie title Pets\n  \"Dogs\" : 3\n```\n")
    assert "diagram not drawn" in deck.warnings[0] and "pie" in deck.warnings[0]
    body = shapes(reopen(prs).slides[1], "deck:body")[0]
    assert body.text_frame.text.startswith("pie title Pets")
    assert body.text_frame.paragraphs[0].runs[0].font.name == "Consolas"


def test_short_text_above_a_visual_long_text_beside_it():
    prs = reopen(convert("## T\n\nOne line.\n\n| a | b |\n|---|--:|\n| x | 1 |\n\n"
                         "## U\n\n" + "A long paragraph of words. " * 30 + "\n\n" + FLOW)[0])
    t, u = prs.slides[1], prs.slides[2]
    above, table = shapes(t, "deck:text")[0], shapes(t, "deck:table")[0]
    assert above.top + above.height <= table.top and table.width > above.width * 0.9
    beside, group = shapes(u, "deck:text")[0], shapes(u, "deck:diagram")[0]
    assert group.left + group.width <= beside.left
    assert table.table.cell(1, 1).text_frame.paragraphs[0].alignment == 3            # --: right
    assert layout_problems(prs) == []


def test_a_long_cell_wraps_and_its_row_grows():
    long = "a sentence that goes on " * 6
    prs = reopen(convert(f"## T\n\n| When | What |\n|---|---|\n| 13:30 | {long} |\n| 13:42 | short |\n")[0])
    table = shapes(prs.slides[1], "deck:table")[0].table
    assert table.cell(1, 1).text_frame.text == long.strip()
    assert table.rows[1].height > table.rows[2].height


def test_pictures_keep_their_ratio_and_alt_text(tmp_path):
    (tmp_path / "a.png").write_bytes(png(400, 100))
    prs = reopen(convert("## P\n\n![a chart](a.png)\n", Options(base_dir=tmp_path))[0])
    (pic,) = shapes(prs.slides[1], "deck:picture")
    assert abs(pic.width / pic.height - 4) < 0.01 and pic._element.nvPicPr.cNvPr.get("descr") == "a chart"
    assert pic.width == 2 * 400 * 9525                     # a small image grows to twice its size at 96 dpi
    (tmp_path / "b.png").write_bytes(png(2000, 500))
    prs = reopen(convert("## P\n\n![](b.png)\n", Options(base_dir=tmp_path))[0])
    (pic,) = shapes(prs.slides[1], "deck:picture")
    assert pic.width > prs.slide_width * 0.8                # a large one fills the slide's width


def test_speaker_notes_and_author(monkeypatch):
    monkeypatch.setenv("MARKDOWN2PPTX_AUTHOR", "Env")
    prs = reopen(convert("## S\n\nx\n\n::: notes\nSay it.\n:::\n")[0])
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Say it."
    assert prs.core_properties.author == "Env"
    prs = reopen(convert("---\nauthor: Ann\n---\n## S\nx\n")[0])
    assert prs.core_properties.author == "Ann"


def test_examples_draw_without_overlap_or_overflow():
    for path in EXAMPLES:
        prs, deck = convert(path.read_text(encoding="utf-8"), Options(base_dir=path.parent))
        assert deck.warnings == [] and layout_problems(reopen(prs)) == []
