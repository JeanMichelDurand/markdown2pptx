"""A user's template: its layouts found by type, its size, theme, body placeholder and footer."""
import io

import pytest
from pptx import Presentation

from markdown2pptx import Options, TemplateError, convert
from markdown2pptx.render.template import layouts
from helpers import EXAMPLES, as_potx, layout_problems, reopen, shapes, template

MD = "# Report\n\n## One\n\n### Text\n\n- a\n- b\n\n## Two\n\n### Diagram\n\n```mermaid\nflowchart LR\nA-->B\n```\n"


def test_the_deck_takes_the_template_size_and_layouts_found_by_type():
    prs = reopen(convert(MD, Options(template=template()))[0])
    assert (prs.slide_width, prs.slide_height) == (9144000, 6858000)          # 10 x 7.5 in, not 16:9
    names = [s.slide_layout.name for s in prs.slides]
    assert names == ["Diapositive de titre", "Titre seul", "Titre de section", "Titre et contenu",
                     "Titre de section", "Titre seul"]                       # the sample slide is gone
    assert prs.slides[0].shapes.title.text == "Report"


@pytest.mark.parametrize("types", [True, False], ids=["by type", "by placeholders"])
def test_layouts_without_types_are_found_by_their_placeholders(types):
    lays = layouts(Presentation(io.BytesIO(template(types=types))))
    assert {k: v.name for k, v in lays.items()} == {
        "title": "Diapositive de titre", "slide": "Titre seul", "text": "Titre et contenu",
        "chapter": "Titre de section" if types else "Diapositive de titre"}


def test_a_potx_template_is_read_as_a_deck():
    prs = reopen(convert(MD, Options(template=as_potx(template())))[0])
    assert prs.slide_width == 9144000


def test_text_keeps_the_template_body_place_and_bullets():
    tpl = Presentation(io.BytesIO(template()))
    lay = layouts(tpl)["text"]
    ph = next(p for p in lay.placeholders if p.placeholder_format.idx == 1)
    prs = reopen(convert(MD, Options(template=template()))[0])
    body = shapes(prs.slides[3], "deck:body")[0]
    assert (body.left, body.top, body.width, body.height) == (ph.left, ph.top, ph.width, ph.height)
    assert all(p._p.pPr is None or p._p.pPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is None for p in body.text_frame.paragraphs)


def test_footer_in_the_template_placeholders():
    prs = reopen(convert(MD, Options(template=template()))[0])
    page = shapes(prs.slides[5], "deck:page")[0]
    assert page.is_placeholder and 'type="slidenum"' in page._element.xml and page.text_frame.text == "6"


@pytest.mark.parametrize("path", EXAMPLES, ids=lambda p: p.stem)
def test_examples_on_a_template_draw_without_overlap_or_overflow(path):
    prs, deck = convert(path.read_text(encoding="utf-8"), Options(template=template(), base_dir=path.parent))
    assert deck.warnings == [] and layout_problems(reopen(prs)) == []


def test_a_file_that_is_not_a_template():
    with pytest.raises(TemplateError, match="not a PowerPoint .pptx or .potx file$"):
        convert(MD, Options(template=b"not a zip"))
