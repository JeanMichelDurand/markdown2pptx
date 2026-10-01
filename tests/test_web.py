"""The browser version's Python glue (web/index.html), run here under CPython on the example."""
import json
import re
from pathlib import Path

from pptx import Presentation

from helpers import EXAMPLES, png, template

PAGE = Path(__file__).parent.parent / "web" / "index.html"
OPTS = {"title": "", "subtitle": "", "aspect": "16:9", "contents": True, "mermaid": "theme", "render": "mermaid"}


def glue(tmp: Path) -> dict:
    page = PAGE.read_text(encoding="utf-8")
    src = page.split("const GLUE = `", 1)[1].split("`;", 1)[0]
    for name in ("out.pptx", "template.pptx", "images"):                  # Pyodide's /tmp, here a tmp dir
        src = src.replace(f'"/tmp/{name}"', repr(str(tmp / name)))
    (tmp / "images").mkdir(exist_ok=True)
    scope: dict = {}
    exec(src, scope)
    return scope


def test_outline_and_run_on_the_example(tmp_path):
    g = glue(tmp_path)
    text = EXAMPLES[0].read_text(encoding="utf-8")
    d = json.loads(g["outline"](text, json.dumps(OPTS), "talk"))
    assert d["title"] == "Moving the team to self-service reporting" and d["chapters"] == 2 and d["warnings"] == []
    assert {s["kind"] for s in d["slides"]} == {"diagram", "table", "text"}
    assert d["slides"][0]["note"].startswith("140 report requests a month reach the data team · Median wait")
    res = json.loads(g["run"](text, json.dumps(OPTS | {"title": "Mine", "mermaid": "slate"}), "talk"))
    assert res["summary"] == d["summary"]
    assert Presentation(tmp_path / "out.pptx").slides[0].shapes.title.text == "Mine"


def test_images_opened_with_the_markdown_are_found_by_name(tmp_path):
    g = glue(tmp_path)
    (tmp_path / "images" / "c.png").write_bytes(png())
    d = json.loads(g["outline"]("## P\n\n![c](figures/c.png)\n", json.dumps(OPTS), "p"))
    assert [s["kind"] for s in d["slides"]] == ["picture"] and d["warnings"] == []


def test_errors_show_their_message_without_the_module_path():
    js = PAGE.read_text(encoding="utf-8").split("const pyError", 1)[1].split(";", 1)[0]
    pattern = js.split(".replace(/", 1)[1].split("/, ", 1)[0]
    assert re.sub(pattern, "", "markdown2pptx.errors.InputError: the document is empty") == "the document is empty"


def test_template_info_and_run_with_a_template(tmp_path):
    (tmp_path / "template.pptx").write_bytes(template())
    g = glue(tmp_path)
    assert json.loads(g["template_info"]()) == {"title": "Diapositive de titre", "chapter": "Titre de section",
                                                "slide": "Titre seul", "text": "Titre et contenu",
                                                "size": "10.00 x 7.50 in"}
    json.loads(g["run"](EXAMPLES[0].read_text(encoding="utf-8"), json.dumps(OPTS | {"template": True}), "talk"))
    assert Presentation(tmp_path / "out.pptx").slide_width == 9144000
