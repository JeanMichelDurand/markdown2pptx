"""Decks built in memory, a user's template, and a layout check on what is drawn."""
import io
import struct
import zlib
from pathlib import Path

from pptx import Presentation

EXAMPLES = sorted((Path(__file__).parent.parent / "examples").glob("*.md"))


def reopen(prs):
    """The deck as PowerPoint will read it: saved, then opened again."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def png(w=40, h=20) -> bytes:
    """A real w x h PNG."""
    raw = b"".join(b"\x00" + b"\xff\x80\x00" * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


FRENCH = {"Title Slide": "Diapositive de titre", "Section Header": "Titre de section", "Title Only": "Titre seul",
          "Title and Content": "Titre et contenu"}


def template(size=(10.0, 7.5), types=True) -> bytes:
    """A user's template: python-pptx's own, renamed in French, resized, with a sample slide to drop.
    types=False strips PowerPoint's layout types, as some generated templates have none."""
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(size[0]), Inches(size[1])
    for layout in prs.slide_layouts:
        layout.name = FRENCH.get(layout.name, layout.name)
        if not types:
            layout._element.attrib.pop("type", None)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Sample slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def as_potx(pptx: bytes) -> bytes:
    """The same deck, typed as PowerPoint types a .potx."""
    import zipfile
    src, buf = zipfile.ZipFile(io.BytesIO(pptx)), io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentation.main+xml", b"template.main+xml")
            out.writestr(item, data)
    return buf.getvalue()


SLACK = 18288            # 0.02 in: a template's own placeholders may touch by rounding


def layout_problems(prs):
    """Shapes off the slide, and shapes that overlap (rules aside)."""
    W, H = prs.slide_width, prs.slide_height
    out = []
    for n, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
            if min(x, y) < 0 or x + w > W or y + h > H:
                out.append(f"slide {n}: {sh.name} off the slide")
            if sh.has_text_frame and not sh.text_frame.text.strip():
                continue                                    # a rule: a filled rectangle with no text
            boxes.append((sh.name, x, y, x + w, y + h))
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                if min(a[3], b[3]) - max(a[1], b[1]) > SLACK and min(a[4], b[4]) - max(a[2], b[2]) > SLACK:
                    out.append(f"slide {n}: {a[0]} overlaps {b[0]}")
    return out


def shapes(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]
