"""A Deck as a Presentation: title slide, contents, a slide per chapter, then the content slides.

The deck is python-pptx's built-in template (Office theme), resized to 16:9 or 4:3, or the user's
template (template.py), whose layouts, fonts and theme colours it takes. Each slide's title sits in
the layout's title placeholder, so PowerPoint's outline, the accessibility checker and a screen
reader find it. A text slide writes into the layout's body placeholder, with the template's
bullets; a slide with a visual (diagram, table, picture) places it on the title-only layout, its
text above it when short, else in a column on the right.
"""
from __future__ import annotations

import io
import math
import os
from datetime import datetime, timezone

import mermaid2pptx
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.parts.image import Image
from pptx.util import Inches

from ..model import Code, Deck, Mermaid, Picture, Slide, Table
from ..options import AUTHOR_ENV, SLIDE_SIZES, Options
from .mermaid import Diagram
from .style import Page, clip, drop_empty_placeholders, footer, placeholder, rule, text
from .table import add_table
from .template import open_template, page as template_page
from .text import fit_size, height, write_blocks

TITLE = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
BODY = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
CONTENTS_ROWS = 10       # per column on the contents slide
LEAD_SHARE = 0.22        # text up to this share of the body's height goes above the visual


def _new(opts: Options, author: str | None):
    """(presentation, page, {kind: layout}): a user's template keeps its own slide size."""
    prs, lays, branded = open_template(opts.template)
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    cp = prs.core_properties                  # replace the built-in deck's own metadata
    if author is None:
        author = os.environ.get(AUTHOR_ENV, "")
    cp.author = cp.last_modified_by = author
    cp.subject, cp.comments, cp.keywords, cp.category = "", "", "", ""
    cp.created = cp.modified = now
    cp.revision = 1
    if branded:
        return prs, template_page(prs, lays["slide"]), lays
    page = Page.of(SLIDE_SIZES[opts.aspect])
    prs.slide_width, prs.slide_height = page.w, page.h
    return prs, page, lays


def _title(s, page: Page, content: str, size: float = 24, color=None):
    if page.branded:                                 # the template's own colour, weight and alignment
        ph = placeholder(s, TITLE, clip(content, 90), size, page.title, None, None, None, None)
    else:
        ph = placeholder(s, TITLE, clip(content, 90), size, page.title, color or page.colors.accent)
    ph.name = "deck:title"
    return ph


def _named(shape, name):
    if shape is not None:
        shape.name = name


def title_slide(prs, page: Page, lays: dict, deck: Deck):
    s = prs.slides.add_slide(lays["title"])
    m, c = page.margin, page.colors
    if page.branded:                                 # the template's title slide, as designed
        _named(placeholder(s, TITLE, deck.title, None, None, bold=None), "deck:title")
        _named(placeholder(s, (PP_PLACEHOLDER.SUBTITLE,), deck.subtitle, None, None, bold=None), "deck:subtitle")
    else:
        rule(s, m, Inches(2.25), Inches(0.9), color=c.accent)
        placeholder(s, TITLE, deck.title, 40, (m, Inches(2.45), page.inner, Inches(1.5)), c.ink).name = "deck:title"
        placeholder(s, (PP_PLACEHOLDER.SUBTITLE,), deck.subtitle, 18,
                    (m, Inches(4.05), page.inner, Inches(1.2)), c.muted).name = "deck:subtitle"
    drop_empty_placeholders(s)


def contents_slide(prs, page: Page, lays: dict, chapters: list[tuple[str, str, int]]) -> dict:
    """chapters: [(number, title, page)]. Returns {page: [boxes]}, linked once the pages exist."""
    s = prs.slides.add_slide(lays["slide"])
    c = page.colors
    _title(s, page, "Contents", 28, c.ink)
    drop_empty_placeholders(s)
    rule(s, page.margin, page.top - Inches(0.25), Inches(0.8), color=c.accent)
    cols = max(1, math.ceil(len(chapters) / CONTENTS_ROWS))
    per = math.ceil(len(chapters) / cols)
    gap, num_w, page_w = Inches(0.35), Inches(0.5), Inches(0.45)
    col_w = int((page.inner - gap * (cols - 1)) / cols)
    title_w = col_w - num_w - page_w
    top = page.top + Inches(0.05)
    pitch = min(Inches(0.55), int((page.bottom - top) / per))
    chars = int(title_w / 914400 * 72 / 7.5)          # about 7.5 pt a character at 14 pt
    links = {}
    for i, (n, title, pg) in enumerate(chapters):
        x, y = page.margin + (i // per) * (col_w + gap), top + pitch * (i % per)
        row_h = pitch - Inches(0.06)
        boxes = (text(s, x, y, num_w, row_h, n, 14, c.accent, True, anchor=MSO_ANCHOR.MIDDLE),
                 text(s, x + num_w, y, title_w, row_h, clip(title, chars), 14, c.ink, anchor=MSO_ANCHOR.MIDDLE),
                 text(s, x + col_w - page_w, y, page_w, row_h, str(pg), 14, c.muted, align=PP_ALIGN.RIGHT,
                      anchor=MSO_ANCHOR.MIDDLE))
        links[pg] = list(boxes[1:])
        rule(s, x, y + pitch - Inches(0.03), col_w, Inches(0.01), c.grid)
    return links


def link_contents(prs, links: dict):
    """Point each contents row's title and page boxes at its chapter slide. The link is on the box,
    not its text: PowerPoint underlines and recolours linked text whatever the run says."""
    for pg, boxes in links.items():
        for box in boxes:
            box.click_action.target_slide = prs.slides[pg - 1]


def chapter_slide(prs, page: Page, lays: dict, number: int, title: str, footer_text: str):
    s = prs.slides.add_slide(lays["chapter"])
    m, c = page.margin, page.colors
    if page.branded:                                 # the template's section header, as designed
        body = placeholder(s, (PP_PLACEHOLDER.BODY,), f"{number:02d}", None, None, bold=None)
        _named(body, "deck:overline")
        heading = clip(title, 120) if body is not None else f"{number:02d}  {clip(title, 116)}"
        _named(placeholder(s, TITLE, heading, None, None, bold=None), "deck:title")
    else:
        placeholder(s, (PP_PLACEHOLDER.BODY,), f"{number:02d}", 20, (m, Inches(2.3), Inches(3), Inches(0.5)),
                    c.accent, bold=True).name = "deck:overline"
        rule(s, m, Inches(2.95), Inches(0.9), color=c.accent)
        placeholder(s, TITLE, clip(title, 120), 36, (m, Inches(3.15), page.inner, Inches(1.6)),
                    c.ink).name = "deck:title"
    drop_empty_placeholders(s)
    footer(s, page, footer_text, len(prs.slides))


def _picture(s, p: Picture, x, y, w, h):
    px_w, px_h = Image.from_blob(p.blob).size
    natural = 914400 // 96                                         # EMU a pixel, at screen resolution
    scale = min(w / max(px_w, 1), h / max(px_h, 1), 2 * natural)   # a small image grows, twice at most
    pw, ph = int(px_w * scale), int(px_h * scale)
    pic = s.shapes.add_picture(io.BytesIO(p.blob), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
    pic.name = "deck:picture"
    if p.alt:
        pic._element.nvPicPr.cNvPr.set("descr", p.alt)


def _body_size(layout) -> float:
    """The template's level-1 body text size in points, where the fit starts (20 when it says none)."""
    sz = layout.slide_master._element.xpath("./p:txStyles/p:bodyStyle/a:lvl1pPr/a:defRPr/@sz")
    return min(28.0, int(sz[0]) / 100) if sz else 20.0


def _overline(s, page: Page, slide: Slide, footer_text: str) -> str:
    over = slide.chapter if slide.chapter and slide.chapter != slide.title else ""
    if over and page.overline:
        text(s, page.margin, page.title[1] - Inches(0.32), page.inner, Inches(0.28), clip(over, 100).upper(), 10,
             page.colors.teal, True, name="deck:overline")
    elif over:                                       # no room above the template's title: the footer says it
        footer_text = " · ".join(t for t in (footer_text, over) if t)
    return footer_text


def text_slide(prs, page: Page, lays: dict, slide: Slide, warn):
    s = prs.slides.add_slide(lays["text"])
    _title(s, page, slide.title)
    body = next((p for p in s.placeholders if p.placeholder_format.type in BODY
                 and p.placeholder_format.idx != 0), None)
    if slide.blocks:
        if body is None or not page.branded:         # the built-in deck's body follows the page grid
            box = (page.margin, page.top, page.inner, page.bottom - page.top)
            if body is not None:
                body.left, body.top, body.width, body.height = box
            start = 24.0
        else:
            box, start = (body.left, body.top, body.width, body.height), _body_size(lays["text"])
        size, fits = fit_size(slide.blocks, box[2], box[3], start)
        if not fits:
            warn("its text does not fit: split it with a --- line")
        if body is None:
            tf = text(s, *box, "", 12, name="deck:body").text_frame
        else:
            body.name, tf = "deck:body", body.text_frame
        write_blocks(tf, slide.blocks, size, page.colors, placeholder=body is not None)
    drop_empty_placeholders(s)
    return s


def visual_slide(prs, page: Page, lays: dict, slide: Slide, opts: Options, diagram, warn) -> tuple:
    s = prs.slides.add_slide(lays["slide"])
    m, c = page.margin, page.colors
    _title(s, page, slide.title)
    drop_empty_placeholders(s)
    top, bottom, extra = page.top, page.bottom, []
    x, w = m, page.inner
    if slide.blocks:
        lead = height(slide.blocks, page.inner, 16) * 12700
        if lead <= (bottom - top) * LEAD_SHARE:      # a short text: above the visual, full width
            box_h = int(lead + Inches(0.1))
            tf = text(s, m, top, page.inner, box_h, "", 12, name="deck:text").text_frame
            write_blocks(tf, slide.blocks, 16, c, placeholder=False)
            top += box_h + Inches(0.15)
        else:                                        # else a column on the right
            side = page.side
            size, fits = fit_size(slide.blocks, side, bottom - top - Inches(0.2), 14)
            if not fits:
                warn("the text beside its visual does not fit: move some of it to a slide of its own")
            col = m + page.inner - side
            rule(s, col, top + Inches(0.05), Inches(0.5), color=c.grid)
            tf = text(s, col, top + Inches(0.2), side, bottom - top - Inches(0.2), "", 12, name="deck:text").text_frame
            write_blocks(tf, slide.blocks, size, c, placeholder=False)
            w = page.inner - side - Inches(0.35)
    v = slide.visual
    if isinstance(v, Mermaid):
        diagram.place(s, x, top, w, bottom - top)
    elif isinstance(v, Table):
        extra = add_table(s, v, x, top, w, bottom - top, opts.max_rows, opts.max_columns)
    else:
        _picture(s, v, x, top, w, bottom - top)
    return s, extra


def content_slide(prs, page: Page, lays: dict, slide: Slide, opts: Options, footer_text: str, deck: Deck):
    n = len(prs.slides) + 1

    def warn(message: str):
        deck.warnings.append(f"slide {n} ({clip(slide.title, 40)}): {message}")

    diagram = None
    if isinstance(slide.visual, Mermaid):            # a diagram that cannot be drawn shows its source
        try:
            diagram = Diagram(slide.visual.src, opts)
            for w in diagram.warnings:
                warn(f"diagram: {w}")
        except (mermaid2pptx.MermaidError, ValueError) as exc:
            warn(f"diagram not drawn, its source is shown instead: {exc}")
            slide.blocks = slide.blocks + [Code(slide.visual.src, "mermaid")]
            slide.visual = None
    extra = []
    if slide.visual is None:
        s = text_slide(prs, page, lays, slide, warn)
    else:
        s, extra = visual_slide(prs, page, lays, slide, opts, diagram, warn)
    footer_text = _overline(s, page, slide, footer_text)
    footer(s, page, " · ".join([t for t in [footer_text] + extra if t]), len(prs.slides))
    if slide.notes:
        s.notes_slide.notes_text_frame.text = slide.notes.strip()


def build(deck: Deck, opts: Options | None = None):
    opts = opts or Options()
    prs, page, lays = _new(opts, deck.author)
    prs.core_properties.title = deck.title
    chapters = deck.chapters()
    contents = opts.contents and len(chapters) >= 2

    # Pages are known before drawing: title, contents, then a chapter slide at each change.
    pages, pg, last = [], 2 + contents, None
    for sl in deck.slides:
        if sl.chapter and sl.chapter != last:
            last = sl.chapter
            pages.append(pg)
            pg += 1
        pg += 1

    title_slide(prs, page, lays, deck)
    links = contents_slide(prs, page, lays, [(f"{i:02d}", t, p) for i, (t, p) in
                                       enumerate(zip(chapters, pages), 1)]) if contents else {}
    if contents:
        footer(prs.slides[-1], page, deck.title, len(prs.slides))
    last, number = None, 0
    for sl in deck.slides:
        if sl.chapter and sl.chapter != last:
            last, number = sl.chapter, number + 1
            chapter_slide(prs, page, lays, number, sl.chapter, deck.title)
        content_slide(prs, page, lays, sl, opts, deck.title, deck)
    link_contents(prs, links)
    return prs

