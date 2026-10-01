"""Text blocks in a text frame: styled runs, real bullets and numbering, code, and a size that fits.

In a template's body placeholder the bullets, indents and colours per level are the template's;
in a plain text box (beside a visual) they are drawn here. Either way the size is chosen so that
the text fits its box (fit_size), from an estimate: PowerPoint does not report text height.
"""
from __future__ import annotations

import math

from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from ..model import Code, ListBlock, Paragraph, Quote, Run, Subheading, plain
from .style import Colors, paint

CODE_FONT = "Consolas"
SIZES = (24, 22, 20, 18, 16, 14, 13, 12, 11, 10)
INDENT = Inches(0.32)                    # a list level, in a text box
EMU_PER_PT = 12700


def _lines(blocks) -> list[tuple[str, float, float, bool]]:
    """(text, size factor, left indent in points, monospace) for each paragraph drawn."""
    out = []
    for b in blocks:
        if isinstance(b, ListBlock):
            out += [(plain(it.runs), 1.0, 26 + 23 * it.level, False) for it in b.items]
        elif isinstance(b, Code):
            out += [(line or " ", 0.8, 0, True) for line in b.text.split("\n")]
        elif isinstance(b, Subheading):
            out.append((plain(b.runs), 1.1, 0, False))
        elif isinstance(b, Quote):
            out.append((plain(b.runs), 1.0, 18, False))
        else:
            out.append((plain(b.runs), 1.0, 0, False))
    return out


def height(blocks, width: int, size: float) -> float:
    """The text's height in points at this size, in a box `width` EMU wide: about 0.5 em a character."""
    total = 0.0
    for text, k, indent, mono in _lines(blocks):
        s = size * k
        per_line = max(8, (width / EMU_PER_PT - indent) / (s * (0.6 if mono else 0.5)))
        total += max(1, math.ceil(len(text) / per_line)) * s * 1.2 + (0 if mono else s * 0.5)
    return total


def fit_size(blocks, width: int, box_height: int, start: float = 20) -> tuple[float, bool]:
    """The largest size from `start` down that fits, and whether one did (the smallest, if none)."""
    sizes = [s for s in SIZES if s <= start] or [SIZES[-1]]
    for s in sizes:
        if height(blocks, width, s) <= box_height / EMU_PER_PT:
            return s, True
    return sizes[-1], False


def write_blocks(tf, blocks, size: float, colors: Colors, placeholder: bool):
    """Fill the text frame: one paragraph per paragraph, list item, quote and code block."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    body = tf._txBody.find(qn("a:bodyPr"))
    for child in list(body):
        if child.tag in (qn("a:spAutoFit"), qn("a:normAutofit"), qn("a:noAutofit")):
            body.remove(child)
    body.append(OxmlElement("a:normAutofit"))            # PowerPoint shrinks the text if it is edited longer
    first = [True]

    def para(level: int = 0):
        p = tf.paragraphs[0] if first[0] else tf.add_paragraph()
        first[0] = False
        if level:
            p.level = min(level, 8)
        return p

    for b in blocks:
        if isinstance(b, ListBlock):
            for it in b.items:
                p = para(it.level if placeholder else 0)
                _bullet(p, it, placeholder)
                runs = ([Run("☑ " if it.checked else "☐ ")] if it.checked is not None else []) + it.runs
                _runs(p, runs, size, colors.ink)
                p.space_after = Pt(size * 0.3)
        elif isinstance(b, Code):
            p = para()
            _no_bullet(p)
            for i, line in enumerate(b.text.split("\n")):
                if i:
                    p.add_line_break()
                _runs(p, [Run(line or " ", code=True)], size * 0.8, colors.ink)
            p.space_after = Pt(size * 0.5)
        elif isinstance(b, Subheading):
            p = para()
            _no_bullet(p)
            _runs(p, [Run(r.text, True, r.italic, r.code, r.strike, r.link) for r in b.runs], size * 1.1,
                  colors.accent)
            p.space_after = Pt(size * 0.25)
        elif isinstance(b, Quote):
            p = para()
            _no_bullet(p, Inches(0.25))
            _runs(p, [Run(r.text, r.bold, True, r.code, r.strike, r.link) for r in b.runs], size, colors.muted)
            p.space_after = Pt(size * 0.5)
        elif isinstance(b, Paragraph):
            p = para()
            _no_bullet(p)
            _runs(p, b.runs, size, colors.ink)
            p.space_after = Pt(size * 0.5)


def _runs(p, runs: list[Run], size: float, color):
    for r in runs:
        run = p.add_run()
        run.text = r.text
        f = run.font
        f.size = Pt(size)
        if r.bold:
            f.bold = True
        if r.italic:
            f.italic = True
        if r.code:
            f.name = CODE_FONT
        if r.strike:
            run._r.get_or_add_rPr().set("strike", "sngStrike")
        if r.link and not r.link.startswith("#"):
            run.hyperlink.address = r.link
        elif color is not None:                          # a link keeps the theme's hyperlink colour
            paint(f.color, color)


def _pPr(p):
    return p._p.get_or_add_pPr()


def _clear_bullet(pPr):
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buBlip"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)


def _append_bullet(pPr, el):
    """pPr's children have an order: the bullet goes before tabs, default run props and extensions."""
    for tag in ("a:tabLst", "a:defRPr", "a:extLst"):
        nxt = pPr.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(el)
            return
    pPr.append(el)


def _no_bullet(p, margin: int = 0):
    pPr = _pPr(p)
    pPr.set("marL", str(int(margin)))
    pPr.set("indent", "0")
    _clear_bullet(pPr)
    _append_bullet(pPr, OxmlElement("a:buNone"))


def _bullet(p, it, placeholder: bool):
    """A template placeholder keeps its own bullet per level, numbers aside; a text box gets • – ·.
    A number gets its own hanging indent: a template's bullet levels may leave it none."""
    pPr = _pPr(p)
    if it.number is not None:
        _clear_bullet(pPr)
        num = OxmlElement("a:buAutoNum")
        num.set("type", "arabicPeriod")
        if it.number != 1:
            num.set("startAt", str(it.number))
        _append_bullet(pPr, num)
    elif not placeholder:
        _clear_bullet(pPr)
        char = OxmlElement("a:buChar")
        char.set("char", "•–·"[min(it.level, 2)])
        _append_bullet(pPr, char)
    if not placeholder or it.number is not None:         # bullets in a placeholder keep the template's indents
        pPr.set("marL", str(int(INDENT * (it.level + 1))))
        pPr.set("indent", str(-int(INDENT)))
