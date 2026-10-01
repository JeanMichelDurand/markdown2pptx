"""A ```mermaid block as native shapes: mermaid2pptx draws it, its group is moved onto the slide.

mermaid2pptx draws the diagram at its natural size (12 pt text) on a slide of its own, as one
group of autoshapes and connectors glued to them. The group is copied into the slide's shape tree
with fresh shape ids (the connectors' ends follow), then scaled into its box: the group's own
transform scales the geometry, and the font sizes are scaled by the same ratio. With the "theme"
colours the shapes use the deck's theme colours and fonts, so they follow the template.
"""
from __future__ import annotations

import copy

import mermaid2pptx
from pptx.oxml.ns import qn

from ..options import Options

MAX_ZOOM = 1.6                    # a small diagram grows, but not into a poster
MIN_FONT = 600                    # 6 pt, in hundredths


class Diagram:
    """A converted diagram, ready to place: its group element, natural size (EMU) and warnings."""

    def __init__(self, src: str, opts: Options):
        mopts = mermaid2pptx.Options(color=opts.mermaid_color, render=opts.mermaid_render, fit=False,
                                     template="", author="")
        prs, parsed, _ = mermaid2pptx.convert(src, mopts)
        self.warnings = list(getattr(parsed, "warnings", []))
        self.summary = parsed.summary() if hasattr(parsed, "summary") else ""
        tree = prs.slides[0].shapes._spTree
        self.group = next(el for el in tree if el.tag == qn("p:grpSp"))
        ext = self.group.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:ext"))
        self.width, self.height = int(ext.get("cx")), int(ext.get("cy"))

    def place(self, slide, x: int, y: int, w: int, h: int, name: str = "deck:diagram"):
        """Copy the group onto the slide, scaled to fit (x, y, w, h) and centred in it."""
        el = copy.deepcopy(self.group)
        _renumber(el, slide)
        r = min(w / max(self.width, 1), h / max(self.height, 1), MAX_ZOOM)
        cw, ch = int(self.width * r), int(self.height * r)
        xfrm = el.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        xfrm.find(qn("a:off")).set("x", str(x + (w - cw) // 2))
        xfrm.find(qn("a:off")).set("y", str(y + (h - ch) // 2))
        xfrm.find(qn("a:ext")).set("cx", str(cw))
        xfrm.find(qn("a:ext")).set("cy", str(ch))             # chOff/chExt keep the drawing's own space
        for tag in ("a:rPr", "a:defRPr", "a:endParaRPr"):
            for rpr in el.iter(qn(tag)):
                if rpr.get("sz"):
                    rpr.set("sz", str(max(MIN_FONT, int(round(int(rpr.get("sz")) * r / 50)) * 50)))
        el.find(qn("p:nvGrpSpPr")).find(qn("p:cNvPr")).set("name", name)
        slide.shapes._spTree.insert_element_before(el, "p:extLst")
        return r


def _renumber(el, slide):
    """Shape ids unique on the slide, the connectors' glued ends renumbered with them."""
    tree = slide.shapes._spTree
    used = [int(c.get("id")) for c in tree.iter(qn("p:cNvPr")) if c.get("id", "").isdigit()]
    nxt = max(used, default=1) + 1
    ids = {}
    for c in el.iter(qn("p:cNvPr")):
        ids[c.get("id")] = str(nxt)
        c.set("id", str(nxt))
        nxt += 1
    for tag in ("a:stCxn", "a:endCxn"):
        for cxn in el.iter(qn(tag)):
            if cxn.get("id") in ids:
                cxn.set("id", ids[cxn.get("id")])
