"""A Table as a native PowerPoint table, in the theme's table style (it follows the deck's colours).

Columns keep the Markdown's alignment (`:--`, `:-:`, `--:`); without one, numbers are right-aligned."""
from __future__ import annotations

import math
import re

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..model import Table
from .style import clip

NUMBER = re.compile(r"^[-+−]?[$€£¥]?\s?[\d.,\s]+%?\s?[kMBK]?$|^(nan|NaN|None|NaT|<NA>|…|\.\.\.)?$")
CELL_CHARS = 160                    # longer cells are cut in the middle; shorter ones wrap
LINE = 1.2                          # line height, in font sizes
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def _numeric(col: list[str]) -> bool:
    return any(c.strip() and c.strip() not in ("...", "…") for c in col) and all(NUMBER.match(c.strip()) for c in col)


def _lines(row: list[str], widths: list[int], size: int) -> int:
    """The row's most wrapped cell, in lines: about half an em a character, as text.py estimates."""
    def lines(text, cw):
        per_line = max(1, int((cw - Inches(0.16)) / (Pt(size) * 0.5)))
        return max(1, math.ceil(len(text) / per_line))
    return max(lines(c, cw) for c, cw in zip(row, widths))


def add_table(slide, t: Table, x, y, w, h, max_rows: int, max_columns: int) -> list[str]:
    """Draw the table at the top of the box; returns what was cut, for the footer."""
    extra: list[str] = []
    width = max(1, len(t.header), *(len(r) for r in t.rows))
    header = (t.header + [""] * width)[:width]
    rows = [(r + [""] * width)[:width] for r in t.rows]
    if len(rows) > max_rows:
        extra.append(f"first {max_rows} of {len(rows)} rows shown")
        rows = rows[:max_rows]
    if width > max_columns:
        extra.append(f"first {max_columns} of {width} columns shown")
        header, rows, width = header[:max_columns], [r[:max_columns] for r in rows], max_columns

    has_header = any(header)
    grid = ([header] if has_header else []) + rows
    grid = [[clip(c, CELL_CHARS) for c in r] for r in grid]
    n = len(grid)
    weights = [min(max(max(len(r[j]) for r in grid), 4), 30) for j in range(width)]
    total = sum(weights)
    widths = [int(w * k / total) for k in weights[:-1]]
    widths.append(w - sum(widths))

    # a long cell wraps: each row is as tall as its most wrapped cell; smaller text if the box is short
    row_h = min(Inches(0.42), int(h / n))
    for size in range(16 if n <= 7 else 14 if n <= 10 else 12 if n <= 13 else 11, 8, -1):
        heights = [max(row_h, int(_lines(r, widths, size) * Pt(size) * LINE + Inches(0.06))) for r in grid]
        if sum(heights) <= h:
            break
    shape = slide.shapes.add_table(n, width, x, y, w, sum(heights))
    shape.name = "deck:table"
    table = shape.table
    table.first_row = has_header
    table.horz_banding = True
    for col, cw in zip(table.columns, widths):
        col.width = cw
    align = (list(t.align) + [None] * width)[:width]
    numeric = [_numeric([r[j] for r in rows]) for j in range(width)]
    for i, row in enumerate(table.rows):
        row.height = heights[i]
        for j, cell in enumerate(row.cells):
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN.get(align[j]) or (PP_ALIGN.RIGHT if numeric[j] else PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = grid[i][j]
            r.font.size = Pt(size)
            r.font.bold = has_header and i == 0
    return extra
