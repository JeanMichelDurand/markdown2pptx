"""Colours, the page grid, and the text helpers every slide is drawn with."""
from __future__ import annotations

from dataclasses import dataclass
from typing import Union

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR as T
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER as P
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# A colour-blind-safe categorical order; the first is the accent. Text takes the theme font.
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
OTHER = "8C8B85"
INK, MUTED, GRID = "1F2328", "5B5F66", "D9DCE1"
ACCENT, TEAL = "174D98", "00797A"

# A colour: "RRGGBB", or (theme colour, brightness) so that a template's theme recolours it.
Color = Union[str, tuple]


@dataclass(frozen=True)
class Colors:
    ink: Color = INK
    muted: Color = MUTED
    grid: Color = GRID
    accent: Color = ACCENT
    teal: Color = TEAL
    other: Color = OTHER
    palette: tuple = tuple(PALETTE)


BUILT_IN = Colors()
THEME = Colors(ink=(T.TEXT_1, 0), muted=(T.TEXT_1, 0.35), grid=(T.TEXT_1, 0.8), accent=(T.ACCENT_1, 0),
               teal=(T.ACCENT_2, 0), other=(T.TEXT_1, 0.55),
               palette=tuple((getattr(T, f"ACCENT_{i}"), 0) for i in range(1, 7)))


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def paint(fmt, color: Color):
    """Set a python-pptx ColorFormat (a font's, a fill's, a line's) to this colour."""
    if isinstance(color, str):
        fmt.rgb = rgb(color)
        return
    fmt.theme_color = color[0]
    if color[1]:
        fmt.brightness = color[1]


@dataclass(frozen=True)
class Page:
    """The slide size, the bands every content slide shares (EMU), and the colours.

    title is the content slides' title box (x, y, w, h); body slides fill top..bottom between the
    side margins. branded: the deck is on a user's template, whose title, chapter and footer
    placeholders are used where the template puts them."""
    w: int
    h: int
    margin: int
    inner: int
    title: tuple
    top: int
    bottom: int
    colors: Colors = BUILT_IN
    branded: bool = False

    @classmethod
    def of(cls, size: tuple[float, float]) -> "Page":
        w, h, m = Inches(size[0]), Inches(size[1]), Inches(0.6)
        return cls(w, h, m, w - 2 * m, (m, Inches(0.62), w - 2 * m, Inches(0.62)), Inches(1.65), h - Inches(0.65))

    @property
    def side(self) -> int:
        """The note column's width, right of the body."""
        return Inches(3.4 if self.inner > Inches(10) else 2.8)

    @property
    def overline(self) -> bool:
        """Room above the title for the chapter's name."""
        return self.title[1] >= Inches(0.55)


def write(tf, text: str, size: float | None, color: Color | None = INK, bold: bool | None = False, align=None,
          spacing: float | None = None, bullets: bool = False):
    """One paragraph per line of `text`. size, color or bold None keeps the placeholder's own."""
    tf.clear()
    for i, para in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if spacing:
            p.line_spacing = spacing
        if bullets:
            p.space_after = Pt(size * 0.6)
        r = p.add_run()
        r.text = ("•  " if bullets else "") + para
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if color:
            paint(r.font.color, color)


def text(slide, x, y, w, h, content: str, size: float = 12, color: Color = INK, bold: bool = False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name: str | None = None, **kw):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    write(tf, content, size, color, bold, align, **kw)
    if name:
        box.name = name
    return box


def placeholder(slide, types: tuple, content: str, size: float | None, box: tuple | None,
                color: Color | None = None, bold: bool | None = False, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """Write into the slide's first placeholder of these types, moved to box=(x, y, w, h); a text box
    there when the layout has none. A placeholder keeps the slide's title for the outline and screen
    readers. box None keeps the layout's place, anchor and alignment (a template's own design);
    size, color, bold, anchor and align None keep the layout's own."""
    ph = next((p for p in slide.placeholders if p.placeholder_format.type in types), None)
    if ph is None:
        return text(slide, *box, content, size or 18, color or INK, bool(bold), anchor=anchor) if box else None
    tf = ph.text_frame
    tf.word_wrap = True
    if box is not None:
        ph.left, ph.top, ph.width, ph.height = box   # all four: a partial set zeroes the others
        tf.auto_size = MSO_AUTO_SIZE.NONE
        if anchor is not None:
            tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = 0
    write(tf, content, size, color, bold, align if box is not None else None)
    return ph


def drop_empty_placeholders(slide):
    """A placeholder left empty shows 'Click to add text' in the editor."""
    for ph in list(slide.placeholders):
        if not (ph.has_text_frame and ph.text_frame.text.strip()):
            ph._element.getparent().remove(ph._element)


def rule(slide, x, y, w, h=Pt(3), color: Color = ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    paint(bar.fill.fore_color, color)
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def footer(slide, page: Page, left: str, number):
    if page.branded and _layout_footer(slide, left, number):
        return
    y, h = page.h - Inches(0.45), Inches(0.3)
    text(slide, page.margin, y, page.inner - Inches(1.1), h, left, size=9, color=page.colors.muted,
         name="deck:footer")
    text(slide, page.margin + page.inner - Inches(1), y, Inches(1), h, str(number), size=9,
         color=page.colors.muted, align=PP_ALIGN.RIGHT, name="deck:page")


def _layout_footer(slide, left: str, number) -> bool:
    """The footer and page number in the layout's own placeholders, styled by the template: False
    when the layout has not both. The number is a field, so it follows when slides move."""
    h = slide.part.package.presentation_part.presentation.slide_height
    found = {ph.placeholder_format.type: ph for ph in slide.slide_layout.placeholders
             if ph.placeholder_format.type in (P.FOOTER, P.SLIDE_NUMBER)
             and ph.top is not None and ph.top + ph.height <= h}
    if len(found) < 2:
        return False
    for kind, name in ((P.FOOTER, "deck:footer"), (P.SLIDE_NUMBER, "deck:page")):
        if kind == P.FOOTER and not left:
            continue
        slide.shapes.clone_placeholder(found[kind])
        ph = next(p for p in slide.placeholders if p.placeholder_format.type == kind)
        ph.name = name
        tf = ph.text_frame
        if kind == P.FOOTER:
            tf.text = clip(left, max(20, int(ph.width / 914400 * 13)))    # about 13 characters an inch
            continue
        p = tf.paragraphs[0]
        fld = OxmlElement("a:fld")
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
        fld.set("type", "slidenum")
        t = OxmlElement("a:t")
        t.text = str(number)
        fld.append(t)
        p._p.append(fld)
    return True


def clip(label, n: int = 40) -> str:
    """Keep head and tail: names often differ only at the end (sales_2024 / sales_2025)."""
    label = str(label)
    if len(label) <= n:
        return label
    head = (n - 1) * 2 // 3
    return label[:head] + "…" + label[-(n - 1 - head):]


def short(content: str, limit: int) -> str:
    """The first sentences of a text, cut at a sentence boundary."""
    content = " ".join(str(content or "").split())
    if len(content) <= limit:
        return content
    cut = content[:limit]
    end = max(cut.rfind(". "), cut.rfind("; "))
    return cut[:end + 1] if end > limit // 3 else cut.rsplit(" ", 1)[0] + " …"
